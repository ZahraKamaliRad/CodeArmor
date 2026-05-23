import os
import json
import re
from pathlib import Path
from collections import defaultdict
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from typing import Dict, List, Tuple
import argparse
from decimal import Decimal
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
import numpy as np

class SecurityMetricsConsolidator:
    def __init__(self, input_dir: str, output_dir: str):
        self.input_dir = Path(input_dir)
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
        self.data = []
        self.method_groups = defaultdict(list)
        self.details_data = []
        self.runtime_data = []

    def parse_metrics_file(self, filepath: Path) -> Dict:
        """Parse a SALLM_metrics.txt file and extract metrics."""
        with open(filepath, 'r') as f:
            content = f.read()
        
        metrics = {}
        
        # Extract run info
        dataset_match = re.search(r'dataset:\s+(\S+)', content)
        model_match = re.search(r'model:\s+(\S+)', content)
        
        if dataset_match:
            metrics['dataset'] = dataset_match.group(1)
        if model_match:
            metrics['model'] = model_match.group(1)
        
        # Extract token usage
        prompt_tokens = re.search(r'prompt_tokens:\s+(\d+)', content)
        completion_tokens = re.search(r'completion_tokens:\s+(\d+)', content)
        total_tokens = re.search(r'total_tokens:\s+(\d+)', content)
        
        if prompt_tokens:
            metrics['prompt_tokens'] = int(prompt_tokens.group(1))
        if completion_tokens:
            metrics['completion_tokens'] = int(completion_tokens.group(1))
        if total_tokens:
            metrics['total_tokens'] = int(total_tokens.group(1))
        
        # Extract runtime
        runtime = re.search(r'runtime_seconds:\s+([\d.]+)', content)
        if runtime:
            metrics['runtime_seconds'] = float(runtime.group(1))
        
        # Extract Bandit metrics
        bandit_section = re.search(r'--- Bandit ---(.+?)(?=--- Semgrep ---|$)', content, re.DOTALL)
        if bandit_section:
            bandit_text = bandit_section.group(1)
            
            n_match = re.search(r'items \(N\):\s+(\d+)', bandit_text)
            v_match = re.search(r'total issues \(V\):\s+(\d+)', bandit_text)
            l_match = re.search(r'total LOC \(L\):\s+(\d+)', bandit_text)
            
            # Extract rate and density from the calculation lines
            rate_match = re.search(r'Vulnerability Rate = V/N = \d+/\d+ = ([\d.]+)', bandit_text)
            density_match = re.search(r'Vulnerability Density = V/L = \d+/\d+ = ([\d.]+)', bandit_text)
            
            if n_match:
                metrics['bandit_items'] = int(n_match.group(1))
            if v_match:
                metrics['bandit_issues'] = int(v_match.group(1))
            if l_match:
                metrics['bandit_loc'] = int(l_match.group(1))
            if rate_match:
                metrics['bandit_rate'] = float(rate_match.group(1))
            if density_match:
                metrics['bandit_density'] = float(density_match.group(1))
        
        # Extract Semgrep metrics
        semgrep_section = re.search(r'--- Semgrep ---(.+?)$', content, re.DOTALL)
        if semgrep_section:
            semgrep_text = semgrep_section.group(1)
            
            n_match = re.search(r'items \(N\):\s+(\d+)', semgrep_text)
            v_match = re.search(r'total issues \(V\):\s+(\d+)', semgrep_text)
            l_match = re.search(r'total LOC \(L\):\s+(\d+)', semgrep_text)
            
            # Extract rate and density from the calculation lines
            rate_match = re.search(r'Vulnerability Rate = V/N = \d+/\d+ = ([\d.]+)', semgrep_text)
            density_match = re.search(r'Vulnerability Density = V/L = \d+/\d+ = ([\d.]+)', semgrep_text)
            
            if n_match:
                metrics['semgrep_items'] = int(n_match.group(1))
            if v_match:
                metrics['semgrep_issues'] = int(v_match.group(1))
            if l_match:
                metrics['semgrep_loc'] = int(l_match.group(1))
            if rate_match:
                metrics['semgrep_rate'] = float(rate_match.group(1))
            if density_match:
                metrics['semgrep_density'] = float(density_match.group(1))
                        
        return metrics

    def extract_method_info(self, method_dir_name: str) -> Tuple[str, int]:
        """Extract method name and run number from directory name."""
        # Handle cases like 'direct(2)' - extract run number in parentheses
        paren_match = re.search(r'\((\d+)\)$', method_dir_name)
        if paren_match:
            run_num = int(paren_match.group(1))
            method = method_dir_name[:paren_match.start()]
            return method, run_num
        
        # For all other cases, keep the full name as the method
        # This preserves rci_iter1, self_refine_iter2, etc.
        return method_dir_name, 1
    
    def collect_data(self):
        """Traverse the input directory and collect all metrics."""
        for db_dir in self.input_dir.iterdir():
            if not db_dir.is_dir():
                continue
            
            database = db_dir.name
            
            for model_dir in db_dir.iterdir():
                if not model_dir.is_dir():
                    continue
                
                model = model_dir.name
                
                for method_dir in model_dir.iterdir():
                    if not method_dir.is_dir():
                        continue
                    
                    method_name, run_num = self.extract_method_info(method_dir.name)
                    
                    # Handle special case for one_shot directory structure
                    if method_name == 'one_shot':
                        for sub_method in ['one_shot', 'zero_shot']:
                            metrics_file = method_dir / sub_method / 'rate' / 'vuln_density' / f'{sub_method}.txt'
                            if metrics_file.exists():
                                metrics = self.parse_metrics_file(metrics_file)
                                metrics['database'] = database
                                metrics['model'] = model
                                metrics['method'] = sub_method
                                metrics['run_number'] = run_num
                                metrics['method_dir'] = method_dir.name
                                self.data.append(metrics)
                                self.method_groups[sub_method].append(metrics)
                            
                            # Collect details_result.xlsx
                            details_file = method_dir / sub_method / 'details_result' / 'details_result.xlsx'
                            if details_file.exists():
                                self.collect_details_result(details_file, database, model, sub_method, run_num)
                    else:
                        # Standard structure
                        metrics_file = method_dir / 'rate' / 'vuln_density' / 'SALLM_metrics.txt'
                        if metrics_file.exists():
                            metrics = self.parse_metrics_file(metrics_file)
                            metrics['database'] = database
                            metrics['model'] = model
                            metrics['method'] = method_name
                            metrics['run_number'] = run_num
                            metrics['method_dir'] = method_dir.name
                            self.data.append(metrics)
                            self.method_groups[method_name].append(metrics)
                        
                        # Collect details_result.xlsx
                        details_file = method_dir / 'details_result' / 'details_result.xlsx'
                        if details_file.exists():
                            self.collect_details_result(details_file, database, model, method_name, run_num)

    def collect_details_result(self, filepath: Path, database: str, model: str, method: str, run_num: int):
        """Collect data from details_result.xlsx files."""
        try:
            df = pd.read_excel(filepath)
            df['database'] = database
            df['model'] = model
            df['technique'] = method
            df['run_number'] = run_num
            self.details_data.append(df)
            
            # Extract runtime information
            runtime_cols = ['total_runtime', 'llm_time', 'bandit_time', 'semgrep_time',
                        'avg_runtime_per_task', 'avg_llm_time_per_task', 
                        'avg_bandit_time_per_task', 'avg_semgrep_time_per_task']
            
            runtime_info = {
                'database': database,
                'model': model,
                'method': method,
                'run_number': run_num
            }
            
            for col in runtime_cols:
                if col in df.columns:
                    runtime_info[col] = df[col].iloc[0] if len(df) > 0 else None
            
            # Also extract token information - check all possible token columns
            token_cols = ['total_api_calls', 'prompt_tokens', 'completion_tokens', 
                        'total_tokens', 'avg_api_calls_per_task',
                        'avg_prompt_tokens_per_task', 'avg_completion_tokens_per_task',
                        'avg_total_tokens_per_task']
            for col in token_cols:
                if col in df.columns:
                    runtime_info[col] = df[col].iloc[0] if len(df) > 0 else None
            
            self.runtime_data.append(runtime_info)
            
        except Exception as e:
            print(f"Warning: Could not read {filepath}: {e}")
          
            
    def count_decimals(self, value):
        """Count decimal places in a value."""
        try:
            d = Decimal(str(value))
            return max(0, -d.as_tuple().exponent)
        except:
            return 0
    
    def process_details_data(self):
        """Process details_result.xlsx files similar to the second script."""
        if not self.details_data:
            return None, None
        
        all_data = pd.concat(self.details_data, ignore_index=True)
        
        # Remove run numbers from technique names
        all_data["technique"] = all_data["technique"].astype(str).str.replace(r"\(\d+\)$", "", regex=True)
        
        group_cols = ["database", "model", "technique"]
        summary_rows = []
        
        for _, group in all_data.groupby(group_cols):
            numeric_cols = group.select_dtypes(include="number").columns
            decimal_map = {}
            for col in numeric_cols:
                first_val = group[col].dropna()
                if len(first_val) > 0:
                    decimal_map[col] = self.count_decimals(first_val.iloc[0])
                else:
                    decimal_map[col] = 2
            
            means = group[numeric_cols].mean()
            for col in numeric_cols:
                means[col] = round(means[col], decimal_map[col])
            
            base_row = group.iloc[0].copy()
            for col in numeric_cols:
                base_row[col] = means[col]
            
            summary_rows.append(base_row)
        
        summary_df = pd.DataFrame(summary_rows)
        summary_df = summary_df[all_data.columns]
        summary_df = summary_df.sort_values(by=["database", "model", "technique"]).reset_index(drop=True)
        
        if 'bandit_rate' in summary_df.columns:
            summary_df["bandit_rate"] = summary_df["bandit_rate"].round(2)
        if 'bandit_density' in summary_df.columns:
            summary_df["bandit_density"] = summary_df["bandit_density"].round(2)
        
        # Calculate reduction rates
        reduction_rows = []
        
        for (dataset, model), group in summary_df.groupby(["dataset", "model"]):
            direct = group[group["technique"] == "direct"]
            if direct.empty:
                continue
            
            direct = direct.iloc[0]
            
            d_bandit_rate = direct.get("bandit_rate", 0)
            d_bandit_density = direct.get("bandit_density", 0)
            d_semgrep_rate = direct.get("semgrep_rate", 0)
            d_semgrep_density = direct.get("semgrep_density", 0)
            
            for _, row in group.iterrows():
                def safe_reduction(base, value):
                    if pd.isna(base) or base == 0:
                        return 0
                    return ((base - value) / base) * 100
                
                reduction_rows.append({
                    "dataset": row["dataset"],
                    "model": row["model"],
                    "technique": row["technique"],
                    "bandit_rate(*100)": row.get("bandit_rate", 0) * 100,
                    "bandit_density(*1000)": row.get("bandit_density", 0) * 1000,
                    "semgrep_rate(*100)": row.get("semgrep_rate", 0) * 100,
                    "semgrep_density(*1000)": row.get("semgrep_density", 0) * 1000,
                    "%▼bandit_rate": round(safe_reduction(d_bandit_rate, row.get("bandit_rate", 0)), 2),
                    "%▼bandit_density": round(safe_reduction(d_bandit_density, row.get("bandit_density", 0)), 2),
                    "%▼semgrep_rate": round(safe_reduction(d_semgrep_rate, row.get("semgrep_rate", 0)), 2),
                    "%▼semgrep_density": round(safe_reduction(d_semgrep_density, row.get("semgrep_density", 0)), 2),
                })
        
        reduction_df = pd.DataFrame(reduction_rows)
        
        if not reduction_df.empty:
            reduction_df = reduction_df[
                ["dataset", "model", "technique",
                 "bandit_rate(*100)", "bandit_density(*1000)",
                 "semgrep_rate(*100)", "semgrep_density(*1000)",
                 "%▼bandit_rate", "%▼bandit_density",
                 "%▼semgrep_rate", "%▼semgrep_density"]
            ]
            reduction_df = reduction_df.sort_values(by=["dataset", "model", "technique"]).reset_index(drop=True)
        
        return summary_df, reduction_df
    
    def merge_runtime_data(self):
        """Merge runtime data from details_result.xlsx into self.data."""
        if not self.runtime_data:
            return
        
        runtime_df = pd.DataFrame(self.runtime_data)
        
        for entry in self.data:
            # Find matching runtime entry
            dataset = entry.get('dataset') or entry.get('database')
            if dataset is None:
                print(f"missing dataset: {entry}")
                continue

            match = runtime_df[
                (runtime_df['database'] == dataset) &
                (runtime_df['model'] == entry['model']) &
                (runtime_df['method'] == entry['method']) &
                (runtime_df['run_number'] == entry['run_number'])
            ]
            
            if not match.empty:
                row = match.iloc[0]
                
                # Override runtime_seconds with total_runtime from Excel
                if pd.notna(row.get('total_runtime')):
                    entry['runtime_seconds'] = row['total_runtime']
                
                # Add detailed runtime breakdown
                for col in ['llm_time', 'bandit_time', 'semgrep_time',
                        'avg_runtime_per_task', 'avg_llm_time_per_task',
                        'avg_bandit_time_per_task', 'avg_semgrep_time_per_task']:
                    if col in row and pd.notna(row[col]):
                        entry[col] = row[col]
                
                # Add token information from Excel (override txt file values if present)
                for col in ['total_api_calls', 'prompt_tokens', 'completion_tokens',
                        'total_tokens', 'avg_api_calls_per_task',
                        'avg_prompt_tokens_per_task', 'avg_completion_tokens_per_task',
                        'avg_total_tokens_per_task']:
                    if col in row and pd.notna(row[col]):
                        entry[col] = row[col]
    
    def create_summary_report(self):
        """Create a comprehensive summary report."""
        df = pd.DataFrame(self.data)
        
        if df.empty:
            print("No data collected. Check input directory structure.")
            return
        
        # Save raw data
        df.to_csv(self.output_dir / 'raw_data.csv', index=False)
        
        # Create summary statistics grouped by method
        summary_cols = ['method', 'bandit_rate', 'bandit_density', 'semgrep_rate', 
                       'semgrep_density', 'total_tokens', 'runtime_seconds']
        
        available_cols = [col for col in summary_cols if col in df.columns]
        
        if len(available_cols) > 1:
            summary = df.groupby('method')[available_cols[1:]].agg(['mean', 'std', 'min', 'max', 'count'])
            summary.to_csv(self.output_dir / 'method_summary.csv')
            
            # Create a flattened version for easier reading
            summary_flat = df.groupby('method').agg({
                col: ['mean', 'std'] for col in available_cols[1:]
            })
            summary_flat.columns = ['_'.join(col).strip() for col in summary_flat.columns.values]
            summary_flat.to_csv(self.output_dir / 'method_summary_flat.csv')
    
    def create_comparison_tables(self):
        """Create detailed comparison tables."""
        df = pd.DataFrame(self.data)
        
        if df.empty:
            return
        
        # Vulnerability Rate Comparison
        if 'bandit_rate' in df.columns and 'semgrep_rate' in df.columns:
            rate_comparison = df.pivot_table(
                values=['bandit_rate', 'semgrep_rate'],
                index='method',
                aggfunc=['mean', 'std', 'count']
            )
            rate_comparison.to_csv(self.output_dir / 'vulnerability_rate_comparison.csv')
        
        # Vulnerability Density Comparison
        if 'bandit_density' in df.columns and 'semgrep_density' in df.columns:
            density_comparison = df.pivot_table(
                values=['bandit_density', 'semgrep_density'],
                index='method',
                aggfunc=['mean', 'std', 'count']
            )
            density_comparison.to_csv(self.output_dir / 'vulnerability_density_comparison.csv')
        
        # Token Usage Comparison
        if 'total_tokens' in df.columns:
            token_comparison = df.pivot_table(
                values=['prompt_tokens', 'completion_tokens', 'total_tokens'],
                index='method',
                aggfunc=['mean', 'std']
            )
            token_comparison.to_csv(self.output_dir / 'token_usage_comparison.csv')
        
        # Runtime Comparison - Enhanced with detailed breakdown
        runtime_cols = ['runtime_seconds']
        detailed_runtime_cols = ['llm_time', 'bandit_time', 'semgrep_time',
                                'avg_runtime_per_task', 'avg_llm_time_per_task',
                                'avg_bandit_time_per_task', 'avg_semgrep_time_per_task']
        
        available_runtime_cols = [col for col in runtime_cols + detailed_runtime_cols if col in df.columns]
        
        if available_runtime_cols:
            runtime_comparison = df.pivot_table(
                values=available_runtime_cols,
                index='method',
                aggfunc=['mean', 'std', 'min', 'max']
            )
            runtime_comparison.to_csv(self.output_dir / 'runtime_comparison.csv')
            
            # Create detailed runtime breakdown
            if any(col in df.columns for col in detailed_runtime_cols):
                runtime_breakdown = df.groupby('method')[available_runtime_cols].mean()
                runtime_breakdown.to_csv(self.output_dir / 'runtime_breakdown.csv')
        
        # Detailed per-run breakdown
        detail_cols = ['method', 'run_number', 'bandit_rate', 'bandit_density', 
                      'semgrep_rate', 'semgrep_density', 'total_tokens', 'runtime_seconds']
        
        # Add detailed runtime columns if available
        for col in detailed_runtime_cols:
            if col in df.columns:
                detail_cols.append(col)
        
        available_detail_cols = [col for col in detail_cols if col in df.columns]
        detailed = df[available_detail_cols]
        detailed = detailed.sort_values(['method', 'run_number'])
        detailed.to_csv(self.output_dir / 'detailed_per_run.csv', index=False)
            
    def create_iteration_analysis(self):
        """Analyze methods with iterations (rci_iter, self_refine_iter, planning)."""
        df = pd.DataFrame(self.data)
        
        if df.empty:
            return
        
        # Extract iteration number for iterative methods
        df['base_method'] = df['method'].str.replace(r'_iter\d+$', '', regex=True)
        df['iteration'] = df['method'].str.extract(r'_iter(\d+)$').fillna(0).astype(int)
        
        # Filter iterative methods
        iterative_methods = df[df['iteration'] > 0]
        
        if not iterative_methods.empty:
            # Build aggregation dict dynamically based on available columns
            agg_dict = {}
            
            # Core metrics to check
            potential_metrics = [
                'bandit_rate', 'bandit_density', 'semgrep_rate', 'semgrep_density',
                'total_tokens', 'prompt_tokens', 'completion_tokens',
                'runtime_seconds', 'llm_time', 'bandit_time', 'semgrep_time',
                'avg_runtime_per_task', 'avg_llm_time_per_task',
                'avg_bandit_time_per_task', 'avg_semgrep_time_per_task',
                'avg_total_tokens_per_task'
            ]
            
            # Only add metrics that exist in the dataframe
            for metric in potential_metrics:
                if metric in iterative_methods.columns:
                    agg_dict[metric] = ['mean', 'std']
            
            # Only proceed if we have metrics to aggregate
            if agg_dict:
                iter_analysis = iterative_methods.groupby(['base_method', 'iteration']).agg(agg_dict)
                iter_analysis.to_csv(self.output_dir / 'iteration_analysis.csv')
                print(f"Created iteration analysis with {len(agg_dict)} metrics")
            else:
                print("Warning: No metrics available for iteration analysis")
                
    def create_runtime_analysis(self):
        """Create comprehensive runtime analysis reports and visualizations."""
        df = pd.DataFrame(self.data)
        
        if df.empty:
            return
        
        # 1. Overall runtime summary by method
        if 'runtime_seconds' in df.columns:
            runtime_summary = df.groupby('method').agg({
                'runtime_seconds': ['mean', 'std', 'min', 'max', 'median']
            })
            runtime_summary.columns = ['mean', 'std', 'min', 'max', 'median']
            runtime_summary = runtime_summary.sort_values('mean', ascending=False)
            runtime_summary.to_csv(self.output_dir / 'runtime_summary.csv')
        
        # 2. Runtime component breakdown
        component_cols = ['llm_time', 'bandit_time', 'semgrep_time']
        available_components = [col for col in component_cols if col in df.columns]
        
        if available_components:
            component_summary = df.groupby('method')[available_components].mean()
            component_summary['total'] = component_summary.sum(axis=1)
            
            # Calculate percentages
            for col in available_components:
                component_summary[f'{col}_pct'] = (component_summary[col] / component_summary['total'] * 100).round(2)
            
            component_summary.to_csv(self.output_dir / 'runtime_component_breakdown.csv')
        
        # 3. Per-task runtime analysis
        per_task_cols = ['avg_runtime_per_task', 'avg_llm_time_per_task', 
                        'avg_bandit_time_per_task', 'avg_semgrep_time_per_task']
        available_per_task = [col for col in per_task_cols if col in df.columns]
        
        if available_per_task:
            per_task_summary = df.groupby('method')[available_per_task].mean()
            per_task_summary.to_csv(self.output_dir / 'runtime_per_task_analysis.csv')
        
        # 4. Runtime efficiency metrics
        if 'runtime_seconds' in df.columns and 'total_tokens' in df.columns:
            efficiency = df.groupby('method').agg({
                'runtime_seconds': 'mean',
                'total_tokens': 'mean'
            })
            efficiency['tokens_per_second'] = (efficiency['total_tokens'] / efficiency['runtime_seconds']).round(2)
            efficiency['seconds_per_1k_tokens'] = (efficiency['runtime_seconds'] / (efficiency['total_tokens'] / 1000)).round(2)
            efficiency.to_csv(self.output_dir / 'runtime_efficiency_metrics.csv')
    def autosize_columns(self, ws):
        """Auto-size columns in Excel worksheet."""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def create_visualizations(self):
        """Create comprehensive visualizations."""
        df = pd.DataFrame(self.data)
        
        if df.empty:
            return
        
        # Set style
        sns.set_style("whitegrid")
        plt.rcParams['figure.figsize'] = (14, 8)
        
        # Identify iterative vs single-run methods
        df['base_method'] = df['method'].str.replace(r'_iter\d+$', '', regex=True)
        df['iteration'] = df['method'].str.extract(r'_iter(\d+)$').fillna(0).astype(int)
        df['method_type'] = df['iteration'].apply(lambda x: 'Iterative' if x > 0 else 'Single-run')
        
        # 1. Vulnerability Density Comparison (VD only, no VR)
        if 'bandit_density' in df.columns and 'semgrep_density' in df.columns:
            fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 6))
            
            density_data = df.groupby('method')[['bandit_density', 'semgrep_density']].mean()
            density_data.plot(kind='bar', ax=ax1, color=['#e74c3c', '#3498db'])
            ax1.set_title('Average Vulnerability Density by Method', fontsize=14, fontweight='bold')
            ax1.set_xlabel('Method', fontsize=12)
            ax1.set_ylabel('Vulnerability Density (issues/LOC)', fontsize=12)
            ax1.legend(['Bandit', 'Semgrep'])
            ax1.tick_params(axis='x', rotation=45)
            ax1.grid(axis='y', alpha=0.3)
            
            # Combined density comparison
            combined_density = df.groupby('method')[['bandit_density', 'semgrep_density']].mean().mean(axis=1).sort_values()
            combined_density.plot(kind='barh', ax=ax2, color='steelblue')
            ax2.set_title('Average Combined Vulnerability Density', fontsize=14, fontweight='bold')
            ax2.set_xlabel('Vulnerability Density (issues/LOC)', fontsize=12)
            ax2.set_ylabel('Method', fontsize=12)
            ax2.grid(axis='x', alpha=0.3)
            
            plt.tight_layout()
            plt.savefig(self.output_dir / 'vulnerability_density_comparison.png', dpi=300, bbox_inches='tight')
            plt.close()
        
        # 2. Box Plot: Single-run Methods VD Distribution
        if 'bandit_density' in df.columns and 'semgrep_density' in df.columns:
            single_run_df = df[df['method_type'] == 'Single-run'].copy()
            
            if not single_run_df.empty:
                fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 7))
                
                # Bandit Density Box Plot
                single_run_df.boxplot(column='bandit_density', by='method', ax=ax1, 
                                    patch_artist=True, showmeans=True,
                                    boxprops=dict(facecolor='lightcoral', alpha=0.7),
                                    medianprops=dict(color='darkred', linewidth=2),
                                    meanprops=dict(marker='D', markerfacecolor='red', markersize=8))
                ax1.set_title('Bandit Vulnerability Density - Single-run Methods', fontsize=13, fontweight='bold')
                ax1.set_xlabel('Method', fontsize=11)
                ax1.set_ylabel('Bandit Density (issues/LOC)', fontsize=11)
                ax1.get_figure().suptitle('')  # Remove automatic title
                ax1.tick_params(axis='x', rotation=45)
                ax1.grid(axis='y', alpha=0.3)
                
                # Semgrep Density Box Plot
                single_run_df.boxplot(column='semgrep_density', by='method', ax=ax2,
                                    patch_artist=True, showmeans=True,
                                    boxprops=dict(facecolor='lightblue', alpha=0.7),
                                    medianprops=dict(color='darkblue', linewidth=2),
                                    meanprops=dict(marker='D', markerfacecolor='blue', markersize=8))
                ax2.set_title('Semgrep Vulnerability Density - Single-run Methods', fontsize=13, fontweight='bold')
                ax2.set_xlabel('Method', fontsize=11)
                ax2.set_ylabel('Semgrep Density (issues/LOC)', fontsize=11)
                ax2.get_figure().suptitle('')
                ax2.tick_params(axis='x', rotation=45)
                ax2.grid(axis='y', alpha=0.3)
                
                plt.tight_layout()
                plt.savefig(self.output_dir / 'boxplot_single_run_vd.png', dpi=300, bbox_inches='tight')
                plt.close()
        
        # 3. Box Plot: Iterative Methods VD Distribution
        if 'bandit_density' in df.columns and 'semgrep_density' in df.columns:
            iterative_df = df[df['method_type'] == 'Iterative'].copy()
            
            if not iterative_df.empty:
                fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 7))
                
                # Bandit Density Box Plot
                iterative_df.boxplot(column='bandit_density', by='method', ax=ax1,
                                    patch_artist=True, showmeans=True,
                                    boxprops=dict(facecolor='#ffcccc', alpha=0.7),
                                    medianprops=dict(color='darkred', linewidth=2),
                                    meanprops=dict(marker='D', markerfacecolor='red', markersize=8))
                ax1.set_title('Bandit Vulnerability Density - Iterative Methods', fontsize=13, fontweight='bold')
                ax1.set_xlabel('Method', fontsize=11)
                ax1.set_ylabel('Bandit Density (issues/LOC)', fontsize=11)
                ax1.get_figure().suptitle('')
                ax1.tick_params(axis='x', rotation=45)
                ax1.grid(axis='y', alpha=0.3)
                
                # Semgrep Density Box Plot
                iterative_df.boxplot(column='semgrep_density', by='method', ax=ax2,
                                    patch_artist=True, showmeans=True,
                                    boxprops=dict(facecolor='#cce5ff', alpha=0.7),
                                    medianprops=dict(color='darkblue', linewidth=2),
                                    meanprops=dict(marker='D', markerfacecolor='blue', markersize=8))
                ax2.set_title('Semgrep Vulnerability Density - Iterative Methods', fontsize=13, fontweight='bold')
                ax2.set_xlabel('Method', fontsize=11)
                ax2.set_ylabel('Semgrep Density (issues/LOC)', fontsize=11)
                ax2.get_figure().suptitle('')
                ax2.tick_params(axis='x', rotation=45)
                ax2.grid(axis='y', alpha=0.3)
                
                plt.tight_layout()
                plt.savefig(self.output_dir / 'boxplot_iterative_vd.png', dpi=300, bbox_inches='tight')
                plt.close()
        
        # 4. Combined Box Plot: All Methods with Color Coding (Iterative vs Single-run)
        if 'bandit_density' in df.columns and 'semgrep_density' in df.columns:
            fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(18, 8))
            
            # Prepare data for box plots with color coding
            methods_sorted = df.groupby('method')['semgrep_density'].median().sort_values().index.tolist()
            
            # Bandit Density - All Methods
            positions = []
            box_data_bandit = []
            colors_bandit = []
            
            for i, method in enumerate(methods_sorted):
                method_data = df[df['method'] == method]
                box_data_bandit.append(method_data['bandit_density'].values)
                positions.append(i)
                # Color by method type
                if method_data['method_type'].iloc[0] == 'Iterative':
                    colors_bandit.append('#ff6b6b')  # Red for iterative
                else:
                    colors_bandit.append('#4ecdc4')  # Teal for single-run
            
            bp1 = ax1.boxplot(box_data_bandit, positions=positions, patch_artist=True,
                            showmeans=True, widths=0.6,
                            medianprops=dict(color='black', linewidth=2),
                            meanprops=dict(marker='D', markerfacecolor='yellow', 
                                        markeredgecolor='black', markersize=6))
            
            for patch, color in zip(bp1['boxes'], colors_bandit):
                patch.set_facecolor(color)
                patch.set_alpha(0.7)
            
            ax1.set_xticks(positions)
            ax1.set_xticklabels(methods_sorted, rotation=45, ha='right')
            ax1.set_title('Bandit Vulnerability Density: All Methods\n(Red=Iterative, Teal=Single-run)', 
                        fontsize=13, fontweight='bold')
            ax1.set_ylabel('Bandit Density (issues/LOC)', fontsize=11)
            ax1.grid(axis='y', alpha=0.3)
            
            # Semgrep Density - All Methods
            box_data_semgrep = []
            colors_semgrep = []
            
            for method in methods_sorted:
                method_data = df[df['method'] == method]
                box_data_semgrep.append(method_data['semgrep_density'].values)
                if method_data['method_type'].iloc[0] == 'Iterative':
                    colors_semgrep.append('#ff6b6b')
                else:
                    colors_semgrep.append('#4ecdc4')
            
            bp2 = ax2.boxplot(box_data_semgrep, positions=positions, patch_artist=True,
                            showmeans=True, widths=0.6,
                            medianprops=dict(color='black', linewidth=2),
                            meanprops=dict(marker='D', markerfacecolor='yellow',
                                        markeredgecolor='black', markersize=6))
            
            for patch, color in zip(bp2['boxes'], colors_semgrep):
                patch.set_facecolor(color)
                patch.set_alpha(0.7)
            
            ax2.set_xticks(positions)
            ax2.set_xticklabels(methods_sorted, rotation=45, ha='right')
            ax2.set_title('Semgrep Vulnerability Density: All Methods\n(Red=Iterative, Teal=Single-run)', 
                        fontsize=13, fontweight='bold')
            ax2.set_ylabel('Semgrep Density (issues/LOC)', fontsize=11)
            ax2.grid(axis='y', alpha=0.3)
            
            # Add legend
            from matplotlib.patches import Patch
            legend_elements = [
                Patch(facecolor='#ff6b6b', alpha=0.7, label='Iterative Methods'),
                Patch(facecolor='#4ecdc4', alpha=0.7, label='Single-run Methods')
            ]
            fig.legend(handles=legend_elements, loc='upper center', ncol=2, 
                    bbox_to_anchor=(0.5, 0.98), fontsize=11)
            
            plt.tight_layout(rect=(0.0, 0.0, 1.0, 0.96))
            plt.savefig(self.output_dir / 'boxplot_all_methods_comparison.png', dpi=300, bbox_inches='tight')
            plt.close()
        
        # 5. Token Usage Comparison
        if 'total_tokens' in df.columns:
            fig, ax = plt.subplots(figsize=(12, 6))
            
            token_data = df.groupby('method')['total_tokens'].mean().sort_values(ascending=False)
            token_data.plot(kind='bar', ax=ax, color='steelblue')
            ax.set_title('Average Total Tokens by Method', fontsize=14, fontweight='bold')
            ax.set_xlabel('Method', fontsize=12)
            ax.set_ylabel('Total Tokens', fontsize=12)
            ax.tick_params(axis='x', rotation=45)
            ax.grid(axis='y', alpha=0.3)
            
            # Add value labels on bars
            for i, v in enumerate(token_data):
                if pd.notna(v):
                    ax.text(i, v, f'{int(v):,}', ha='center', va='bottom')
            
            plt.tight_layout()
            plt.savefig(self.output_dir / 'token_usage.png', dpi=300, bbox_inches='tight')
            plt.close()
        
        # 6. Runtime Comparison
        if 'runtime_seconds' in df.columns:
            fig, ax = plt.subplots(figsize=(12, 6))
            
            runtime_data = df.groupby('method')['runtime_seconds'].mean().sort_values(ascending=False)
            runtime_data.plot(kind='bar', ax=ax, color='coral')
            ax.set_title('Average Runtime by Method', fontsize=14, fontweight='bold')
            ax.set_xlabel('Method', fontsize=12)
            ax.set_ylabel('Runtime (seconds)', fontsize=12)
            ax.tick_params(axis='x', rotation=45)
            ax.grid(axis='y', alpha=0.3)
            
            # Add value labels on bars
            for i, v in enumerate(runtime_data):
                if pd.notna(v):
                    ax.text(i, v, f'{v:.1f}s', ha='center', va='bottom')
            
            plt.tight_layout()
            plt.savefig(self.output_dir / 'runtime_comparison.png', dpi=300, bbox_inches='tight')
            plt.close()
        
        # 7. Runtime Component Breakdown
        component_cols = ['llm_time', 'bandit_time', 'semgrep_time']
        available_components = [col for col in component_cols if col in df.columns]
        
        if len(available_components) >= 2:
            fig, ax = plt.subplots(figsize=(14, 7))
            
            component_data = df.groupby('method')[available_components].mean()
            component_data.plot(kind='bar', stacked=True, ax=ax, 
                            color=['#3498db', '#e74c3c', '#2ecc71'])
            ax.set_title('Runtime Component Breakdown by Method', fontsize=14, fontweight='bold')
            ax.set_xlabel('Method', fontsize=12)
            ax.set_ylabel('Time (seconds)', fontsize=12)
            ax.legend(title='Component', labels=['LLM Time', 'Bandit Time', 'Semgrep Time'])
            ax.tick_params(axis='x', rotation=45)
            ax.grid(axis='y', alpha=0.3)
            
            plt.tight_layout()
            plt.savefig(self.output_dir / 'runtime_components.png', dpi=300, bbox_inches='tight')
            plt.close()
        
        # 8. Iteration Analysis (VD only)
        iterative_methods = df[df['iteration'] > 0]
        
        if not iterative_methods.empty and 'bandit_density' in df.columns:
            fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 6))
            
            # Bandit Density by Iteration
            for base_method in iterative_methods['base_method'].unique():
                method_data = iterative_methods[iterative_methods['base_method'] == base_method]
                iter_means = method_data.groupby('iteration')['bandit_density'].mean()
                ax1.plot(iter_means.index, iter_means.values, marker='o', label=base_method, linewidth=2)
            
            ax1.set_title('Bandit Vulnerability Density by Iteration', fontsize=12, fontweight='bold')
            ax1.set_xlabel('Iteration', fontsize=10)
            ax1.set_ylabel('Bandit Density (issues/LOC)', fontsize=10)
            ax1.legend()
            ax1.grid(True, alpha=0.3)
            
            # Semgrep Density by Iteration
            if 'semgrep_density' in iterative_methods.columns:
                for base_method in iterative_methods['base_method'].unique():
                    method_data = iterative_methods[iterative_methods['base_method'] == base_method]
                    iter_means = method_data.groupby('iteration')['semgrep_density'].mean()
                    ax2.plot(iter_means.index, iter_means.values, marker='o', label=base_method, linewidth=2)
                
                ax2.set_title('Semgrep Vulnerability Density by Iteration', fontsize=12, fontweight='bold')
                ax2.set_xlabel('Iteration', fontsize=10)
                ax2.set_ylabel('Semgrep Density (issues/LOC)', fontsize=10)
                ax2.legend()
                ax2.grid(True, alpha=0.3)
            
            plt.tight_layout()
            plt.savefig(self.output_dir / 'iteration_density_analysis.png', dpi=300, bbox_inches='tight')
            plt.close()
        
        # 9. Efficiency Scatter Plot (VD only)
        if 'runtime_seconds' in df.columns and 'total_tokens' in df.columns and 'semgrep_density' in df.columns:
            fig, ax = plt.subplots(figsize=(12, 8))
            
            method_summary = df.groupby('method').agg({
                'runtime_seconds': 'mean',
                'total_tokens': 'mean',
                'semgrep_density': 'mean'
            })
            
            scatter = ax.scatter(method_summary['runtime_seconds'], 
                            method_summary['total_tokens'],
                            s=method_summary['semgrep_density'] * 5000,
                            alpha=0.6,
                            c=range(len(method_summary)),
                            cmap='viridis')
            
            for idx, method in enumerate(method_summary.index):
                ax.annotate(method, 
                        (method_summary.loc[method, 'runtime_seconds'],
                        method_summary.loc[method, 'total_tokens']),
                        xytext=(5, 5), textcoords='offset points', fontsize=9)
            
            ax.set_xlabel('Runtime (seconds)', fontsize=12)
            ax.set_ylabel('Total Tokens', fontsize=12)
            ax.set_title('Method Efficiency: Runtime vs Token Usage\n(Bubble size = Semgrep Density)', 
                        fontsize=14, fontweight='bold')
            ax.grid(True, alpha=0.3)
            
            plt.tight_layout()
            plt.savefig(self.output_dir / 'efficiency_scatter.png', dpi=300, bbox_inches='tight')
            plt.close()

    def create_excel_report(self):
        """Create a comprehensive Excel report with multiple sheets."""
        df = pd.DataFrame(self.data)
        
        if df.empty:
            return
        
        excel_path = self.output_dir / 'comprehensive_report.xlsx'
        
        with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
            # Sheet 1: Raw Data
            df.to_excel(writer, sheet_name='Raw Data', index=False)
            
            # Sheet 2: Method Summary
            summary_cols = ['method', 'bandit_rate', 'bandit_density', 'semgrep_rate', 
                           'semgrep_density', 'total_tokens', 'runtime_seconds']
            available_cols = [col for col in summary_cols if col in df.columns]
            
            if len(available_cols) > 1:
                summary = df.groupby('method')[available_cols[1:]].agg(['mean', 'std', 'min', 'max'])
                summary.to_excel(writer, sheet_name='Method Summary')
            
            # Sheet 3: Vulnerability Rates
            if 'bandit_rate' in df.columns and 'semgrep_rate' in df.columns:
                rate_comparison = df.pivot_table(
                    values=['bandit_rate', 'semgrep_rate'],
                    index='method',
                    aggfunc=['mean', 'std', 'count']
                )
                rate_comparison.to_excel(writer, sheet_name='Vulnerability Rates')
            
            # Sheet 4: Token Usage
            if 'total_tokens' in df.columns:
                token_comparison = df.pivot_table(
                    values=['prompt_tokens', 'completion_tokens', 'total_tokens'],
                    index='method',
                    aggfunc=['mean', 'std']
                )
                token_comparison.to_excel(writer, sheet_name='Token Usage')
            
            # Sheet 5: Runtime Analysis
            if 'runtime_seconds' in df.columns:
                runtime_comparison = df.pivot_table(
                    values=['runtime_seconds'],
                    index='method',
                    aggfunc=['mean', 'std', 'min', 'max']
                )
                runtime_comparison.to_excel(writer, sheet_name='Runtime Analysis')
            
            # Sheet 6: Runtime Component Breakdown
            component_cols = ['llm_time', 'bandit_time', 'semgrep_time']
            available_components = [col for col in component_cols if col in df.columns]
            
            if available_components:
                component_summary = df.groupby('method')[available_components].mean()
                component_summary['total'] = component_summary.sum(axis=1)
                
                for col in available_components:
                    component_summary[f'{col}_pct'] = (component_summary[col] / component_summary['total'] * 100).round(2)
                
                component_summary.to_excel(writer, sheet_name='Runtime Components')
            
            # Sheet 7: Efficiency Metrics
            if 'runtime_seconds' in df.columns and 'total_tokens' in df.columns:
                efficiency = df.groupby('method').agg({
                    'runtime_seconds': 'mean',
                    'total_tokens': 'mean'
                })
                efficiency['tokens_per_second'] = (efficiency['total_tokens'] / efficiency['runtime_seconds']).round(2)
                efficiency['seconds_per_1k_tokens'] = (efficiency['runtime_seconds'] / (efficiency['total_tokens'] / 1000)).round(2)
                efficiency.to_excel(writer, sheet_name='Efficiency Metrics')
            
            # Sheet 8: Details Summary (from details_result.xlsx)
            summary_df, reduction_df = self.process_details_data()
            if summary_df is not None:
                summary_df.to_excel(writer, sheet_name='Details Summary', index=False)
            
            # Sheet 9: Reduction Rates
            if reduction_df is not None and not reduction_df.empty:
                reduction_df.to_excel(writer, sheet_name='Reduction Rates', index=False)
            
            # Sheet 10: Iteration Analysis
            df['base_method'] = df['method'].str.replace(r'_iter\d+$', '', regex=True)
            df['iteration'] = df['method'].str.extract(r'_iter(\d+)$').fillna(0).astype(int)
            iterative_methods = df[df['iteration'] > 0]
            
            if not iterative_methods.empty:
                agg_dict = {
                    'bandit_rate': ['mean', 'std'],
                    'bandit_density': ['mean', 'std'],
                    'semgrep_rate': ['mean', 'std'],
                    'semgrep_density': ['mean', 'std'],
                    'total_tokens': ['mean', 'std'],
                    'runtime_seconds': ['mean', 'std']
                }
                
                iter_analysis = iterative_methods.groupby(['base_method', 'iteration']).agg(agg_dict)
                iter_analysis.to_excel(writer, sheet_name='Iteration Analysis')
        
        # Auto-size columns
        wb = load_workbook(excel_path)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            self.autosize_columns(ws)
            
            # Center align headers
            for cell in ws[1]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
        
        wb.save(excel_path)
        print(f"Comprehensive Excel report saved to: {excel_path}")
    
    def run(self):
        """Execute the complete analysis pipeline."""
        print("Starting data collection...")
        self.collect_data()
        
        print(f"Collected {len(self.data)} data points from {len(self.method_groups)} methods")
        
        if not self.data:
            print("No data found. Please check the input directory structure.")
            return
        
        print("Merging runtime data from details_result.xlsx...")
        self.merge_runtime_data()
        
        print("Creating summary reports...")
        self.create_summary_report()
        
        print("Creating comparison tables...")
        self.create_comparison_tables()
        
        print("Creating iteration analysis...")
        self.create_iteration_analysis()
        
        print("Creating runtime analysis...")
        self.create_runtime_analysis()
        
        print("Creating visualizations...")
        self.create_visualizations()
        
        print("Creating comprehensive Excel report...")
        self.create_excel_report()
        
        print(f"\nAnalysis complete! Results saved to: {self.output_dir}")
        print("\nGenerated files:")
        for file in sorted(self.output_dir.iterdir()):
            print(f"  - {file.name}")


def main():
    parser = argparse.ArgumentParser(
        description='Consolidate and analyze security metrics from SALLM experiments'
    )
    parser.add_argument(
        '--input-dir',
        type=str,
        required=True,
        help='Input directory containing the experiment results'
    )
    parser.add_argument(
        '--output-dir',
        type=str,
        default='consolidated_results',
        help='Output directory for consolidated results (default: consolidated_results)'
    )
    
    args = parser.parse_args()
    
    consolidator = SecurityMetricsConsolidator(args.input_dir, args.output_dir)
    consolidator.run()


if __name__ == '__main__':
    main()

# python .\scripts\results.py --input-dir .\outputs\ --output-dir .\reports\ 