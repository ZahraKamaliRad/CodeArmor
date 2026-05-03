from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from .metrics import compute_metrics

def autosize_columns(ws):
    for column in ws.columns:
        max_length = 0
        col = column[0].column
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col)].width = max_length + 2

def center_align(ws):
    align = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = align

def save_experiment_summary(
    out_dir: str,
    dataset: str,
    model: str,
    technique: str,
    total_tasks: int,
    elapsed_time: float,
    total_llm_time: float,
    total_bandit_time: float,
    total_semgrep_time: float,
    total_api_calls: int,
    total_prompt_tokens: int,
    total_completion_tokens: int,
    jsonl_path: str,
):

    metrics = compute_metrics(jsonl_path)

    bandit_rate = metrics.get("bandit", {}).get("vuln_rate_value", 0)
    bandit_density = metrics.get("bandit", {}).get("density_value", 0)

    semgrep_rate = metrics.get("semgrep", {}).get("vuln_rate_value", 0)
    semgrep_density = metrics.get("semgrep", {}).get("density_value", 0)

    avg_runtime_per_task = round(elapsed_time / total_tasks, 4) if total_tasks else 0
    avg_llm_time_per_task = round(total_llm_time / total_tasks, 4) if total_tasks else 0
    avg_bandit_time_per_task = round(total_bandit_time / total_tasks, 4) if total_tasks else 0
    avg_semgrep_time_per_task = round(total_semgrep_time / total_tasks, 4) if total_tasks else 0
    avg_api_calls_per_task = round(total_api_calls / total_tasks, 4) if total_tasks else 0

    out_dir = Path(out_dir)
    detail_dir = out_dir / "details_result"
    detail_dir.mkdir(parents=True, exist_ok=True)
    xlsx_file = detail_dir / f"{detail_dir.name}.xlsx"

    headers = [
        "dataset",
        "model",
        "technique",
        "total_tasks",
        "bandit_rate",
        "bandit_density",
        "semgrep_rate",
        "semgrep_density",
        "total_runtime",
        "llm_time",
        "bandit_time",
        "semgrep_time",
        "total_api_calls",
        "prompt_tokens",
        "completion_tokens",
        "total_tokens",
        "avg_runtime_per_task",
        "avg_llm_time_per_task",
        "avg_bandit_time_per_task",
        "avg_semgrep_time_per_task",
        "avg_api_calls_per_task",
    ]

    row = [
        dataset,
        model,
        technique,
        total_tasks,
        bandit_rate,
        bandit_density,
        semgrep_rate,
        semgrep_density,
        round(elapsed_time, 4),
        round(total_llm_time, 4),
        round(total_bandit_time, 4),
        round(total_semgrep_time, 4),
        total_api_calls,
        total_prompt_tokens,
        total_completion_tokens,
        total_prompt_tokens + total_completion_tokens,
        avg_runtime_per_task,
        avg_llm_time_per_task,
        avg_bandit_time_per_task,
        avg_semgrep_time_per_task,
        avg_api_calls_per_task,
    ]

    try:
        wb = Workbook()
        ws = wb.active

        ws.append(headers)
        ws.append(row)

        autosize_columns(ws)
        center_align(ws)

        wb.save(xlsx_file)

        print(f"Summary metrics saved to: {xlsx_file}")

    except Exception as e:
        print(f"Error saving summary metrics: {e}")
