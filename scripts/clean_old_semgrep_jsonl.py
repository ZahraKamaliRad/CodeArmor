from __future__ import annotations
import json
import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
from datetime import datetime
import sys
import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from src.utils.metrics import save_metrics_result
from src.utils.plot_refinment import plot_totals
from src.utils.metrics import compute_metrics


SECURITY_RULE_PATTERNS = [
    "security","injection","sql","xss","ssti","ssrf","csrf","deserialization",
    "command","exec","eval","path-traversal","traversal","open-redirect","redirect",
    "crypto","hash","weak-random","random","jwt","auth","authentication",
    "authorization","secret","hardcoded","token","password","pickle","xml",
    "xxe","ldap","rce","dos","overflow","taint","unsafe",
]


def find_project_root() -> Path:
    current = Path(__file__).resolve()

    for parent in [current.parent] + list(current.parents):
        if (parent / "outputs").exists():
            return parent

    cwd = Path.cwd().resolve()
    if (cwd / "outputs").exists():
        return cwd

    raise RuntimeError("Could not find project root containing 'outputs'.")


def backup_outputs(outputs_dir: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = outputs_dir / f"_backup_before_semgrep_cleanup_{timestamp}"

    if backup_dir.exists():
        raise RuntimeError("Backup directory already exists")

    print(f"[INFO] Creating backup: {backup_dir}")

    shutil.copytree(
        outputs_dir,
        backup_dir,
        ignore=shutil.ignore_patterns("_backup_before_semgrep_cleanup*")
    )

    print("[INFO] Backup completed")

    return backup_dir


def is_real_security_issue(rule_id: Optional[str],message: str,
    metadata: Dict[str, Any],cwe: Optional[Union[str, List[str]]]) -> bool:

    if cwe:
        return True

    text = " ".join([str(rule_id or ""), str(message or "")]).lower()
    category = str(metadata.get("category") or "").lower()
    confidence = str(metadata.get("confidence") or "").lower()
    impact = str(metadata.get("impact") or "").lower()

    if category == "security":
        return True

    for pattern in SECURITY_RULE_PATTERNS:
        if pattern in text:
            return True

    meta_text = " ".join([category, confidence, impact]).lower()

    for pattern in SECURITY_RULE_PATTERNS:
        if pattern in meta_text:
            return True

    return False


def clean_semgrep_block(block: Dict[str, Any]) -> tuple[int, int]:

    issues = block.get("issues", [])
    if not isinstance(issues, list):
        issues = []

    before = len(issues)
    filtered: List[Dict[str, Any]] = []

    for issue in issues:
        if not isinstance(issue, dict):
            continue

        if is_real_security_issue(issue.get("rule_id"),str(issue.get("message") or ""),
            issue.get("metadata") or {},issue.get("cwe")):
            filtered.append(issue)

    block["issues"] = filtered
    block["secure"] = len(filtered) == 0

    severity_counts = {}
    cwe_counts = {}

    for issue in filtered:
        severity = str(issue.get("severity") or "unknown")
        severity_counts[severity] = severity_counts.get(severity, 0) + 1

        cwe = issue.get("cwe")
        if isinstance(cwe, list):
            for c in cwe:
                cwe_counts[str(c)] = cwe_counts.get(str(c), 0) + 1
        elif cwe:
            cwe_counts[str(cwe)] = cwe_counts.get(str(cwe), 0) + 1

    block["summary"] = {"total_tool_issues": len(filtered),"severity_counts": severity_counts,
        "cwe_counts": cwe_counts,"secure": len(filtered) == 0}

    return before, len(filtered)


def process_record(record: Dict[str, Any]) -> tuple[int, int]:

    total_before = 0
    total_after = 0

    if isinstance(record.get("semgrep_result"), dict):
        b, a = clean_semgrep_block(record["semgrep_result"])
        total_before += b
        total_after += a

    if isinstance(record.get("initial_analysis"), dict):
        ia = record["initial_analysis"]
        if isinstance(ia.get("semgrep_result"), dict):
            b, a = clean_semgrep_block(ia["semgrep_result"])
            total_before += b
            total_after += a

    for it in record.get("iterations", []):
        if not isinstance(it, dict):
            continue

        analysis = it.get("analysis", {})
        if isinstance(analysis.get("semgrep_result"), dict):
            b, a = clean_semgrep_block(analysis["semgrep_result"])
            total_before += b
            total_after += a

    return total_before, total_after



def update_semgrep_in_excel(parent_dir: Path, jsonl_path: Path):

    excel_path = parent_dir / "details_result" / "details_result.xlsx"

    if not excel_path.exists():
        return

    metrics = compute_metrics(jsonl_path)

    semgrep_rate = metrics.get("semgrep", {}).get("vuln_rate_value", 0)
    semgrep_density = metrics.get("semgrep", {}).get("density_value", 0)

    wb = load_workbook(excel_path)
    ws = wb.active

    header_map = {}

    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col).value
        if header:
            header_map[header] = col

    if "semgrep_rate" in header_map:
        ws.cell(row=2, column=header_map["semgrep_rate"]).value = semgrep_rate

    if "semgrep_density" in header_map:
        ws.cell(row=2, column=header_map["semgrep_density"]).value = semgrep_density

    wb.save(excel_path)

    print(f"excel semgrep updated → {semgrep_rate}, {semgrep_density}")

def process_jsonl(file_path: Path):

    total_before = 0
    total_after = 0
    lines_out = []

    with file_path.open("r", encoding="utf-8") as f:
        for line in f:
            if not line.strip():
                continue

            record = json.loads(line)

            b, a = process_record(record)

            total_before += b
            total_after += a

            lines_out.append(json.dumps(record, ensure_ascii=False))

    with file_path.open("w", encoding="utf-8") as f:
        for l in lines_out:
            f.write(l + "\n")

    print(f"  semgrep issues: {total_before} -> {total_after}")

    parent = file_path.parent

    save_metrics_result(str(file_path))

    update_semgrep_in_excel(parent, file_path)

    plot_path = parent / "plots" / "Refinement.png"

    plot_totals(jsonl_path=file_path,out=plot_path,show=False,cumulative=True)

    print(f"  plot regenerated: {plot_path}")


def main():

    project_root = find_project_root()
    outputs_dir = project_root / "outputs"

    print(f"[INFO] outputs dir: {outputs_dir}")

    backup_outputs(outputs_dir)

    jsonl_files = []

    for subdir in outputs_dir.iterdir():
        if subdir.is_dir() and subdir.name.startswith("_backup"):
            continue
        jsonl_files.extend(subdir.rglob("*.jsonl"))

    print(f"[INFO] found {len(jsonl_files)} jsonl files")

    for i, path in enumerate(jsonl_files, 1):
        print(f"\n[{i}/{len(jsonl_files)}] Cleaning: {path}")
        try:
            process_jsonl(path)
        except Exception as e:
            print(f"[ERROR] {path}: {e}")

    print("\nDONE")


if __name__ == "__main__":
    main()
