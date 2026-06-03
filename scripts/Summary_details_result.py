
import os
import re
import pandas as pd
from pathlib import Path
from decimal import Decimal, ROUND_DOWN
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


class SecurityDataAggregator:
    def __init__(self):
        self.base_path = Path(os.path.dirname(os.path.abspath(__file__)))
        self.project_root = self.base_path.parent
        self.base_dir = self.project_root / "outputs"
        self.output_file = self.base_path / "summary.xlsx"

        self.TECHNIQUE_SUFFIX_RE = re.compile(r"\(\d+\)$", re.IGNORECASE)

        self.LOC_RE = re.compile(
            r"^\s*total\s+LOC\s+\(L\)\s*:\s*(\d+)\s*$",
            re.IGNORECASE | re.MULTILINE,
        )

        self.PASS1_RE = re.compile(
            r"^\s*Pass@1\s*:\s*([0-9.]+)\s*$",
            re.IGNORECASE | re.MULTILINE,
        )

        self.functionality_dir = self.project_root / "functionality" / "outputs"

        self.TECHNIQUE_MAP = {
            "direct": "Direct",
            "prefix": "Prefix priming",
            "persona": "Persona priming",
            "zero_shot_cot": "Zero-shot CoT",
            "few_shot_cot": "Few-shot CoT",
            "one_shot": "One-shot",
            "self_refine_iter1": "Self-refine (1-iter)",
            "self_refine_iter2": "Self-refine (2-iter)",
            "rci_iter1": "RCI (1-iter)",
            "rci_iter2": "RCI (2-iter)",
            "planning_zero_shot": "planning(basic)",
            "planning_zero_shot_iter1": "planning(1-iter)",
            "planning_zero_shot_iter2": "planning(2-iter)",
        }

    def _truncate4(self, x):
        try:
            return float(
                Decimal(str(x)).quantize(Decimal("0.0001"), rounding=ROUND_DOWN)
            )
        except:
            return x

    def _normalize_technique(self, name):
        if not name:
            return "unknown"
        clean_name = self.TECHNIQUE_SUFFIX_RE.sub("", str(name)).strip().lower()
        return self.TECHNIQUE_MAP.get(clean_name, clean_name)

    def _format_excel(self):
        wb = load_workbook(self.output_file)
        align = Alignment(horizontal="center", vertical="center")

        for ws in wb.worksheets:
            for col in ws.columns:
                max_length = max(
                    (len(str(cell.value)) for cell in col if cell.value), default=0
                )
                ws.column_dimensions[get_column_letter(col[0].column)].width = (
                    max_length + 2
                )

            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = align

        model_mapping = {
            "gemini-flash-lite-latest": "Gemini",
            "gpt-4o-mini": "GPT-4",
            "GPT-OSS-120B-9az8h": "GPT-OSS-120B",
            "Qwen3-30B-A3B-70w4t": "Qwen3-30B-A3B",     
        }

        dataset_mapping = {
            "sallm": "SALLM",
            "securityeval": "SecurityEval",
        }

        for ws in wb.worksheets:
            model_col_idx = None
            dataset_col_idx = None

            for cell in ws[1]:
                col_name = str(cell.value).strip().lower()

                if col_name == "model":
                    model_col_idx = cell.column

                if col_name == "dataset":
                    dataset_col_idx = cell.column

            if model_col_idx:
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row, column=model_col_idx)
                    if cell.value in model_mapping:
                        cell.value = model_mapping[cell.value]

            if dataset_col_idx:
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row, column=dataset_col_idx)
                    if isinstance(cell.value, str):
                        value = cell.value.strip().lower()
                        if value in dataset_mapping:
                            cell.value = dataset_mapping[value]

        wb.save(self.output_file)

    def run_aggregation(self):
        all_dfs = []
        loc_rows = []

        for root, _, files in os.walk(self.base_dir):
            rel_path = os.path.relpath(root, self.base_dir)
            parts = rel_path.split(os.sep)

            for file in files:
                full_path = Path(root) / file

                if file.endswith(".xlsx") and not file.startswith("~$"):
                    try:
                        all_dfs.append(pd.read_excel(full_path))
                    except:
                        pass

                if file.endswith(".txt") and len(parts) >= 3:
                    try:
                        with open(
                            full_path, "r", encoding="utf-8", errors="ignore"
                        ) as f:
                            match = self.LOC_RE.search(f.read())
                            if match:
                                loc_rows.append(
                                    {
                                        "dataset": parts[0],
                                        "model": parts[1],
                                        "technique": self._normalize_technique(
                                            parts[2]
                                        ),
                                        "total LOC (L)": int(match.group(1)),
                                    }
                                )
                    except:
                        pass

        self._process_and_save(all_dfs, loc_rows)

    def _process_and_save(self, all_dfs, loc_rows):
        if not all_dfs:
            return

        master_df = pd.concat(all_dfs, ignore_index=True)
        master_df["technique"] = master_df["technique"].apply(
            self._normalize_technique
        )

        token_cols = ["prompt_tokens", "completion_tokens", "total_tokens"]
        existing_tokens = [c for c in token_cols if c in master_df.columns]
        
        token_summary_df = (
            master_df.groupby(["dataset", "model", "technique"])[existing_tokens]
            .sum()
            .reset_index()
            .sort_values(["dataset", "model", "technique"])
        )

        summary_rows = []
        for _, gp in master_df.groupby(["dataset", "model", "technique"]):
            num_cols = gp.select_dtypes(include="number").columns
            means = gp[num_cols].mean()
            row = gp.iloc[0].copy()
            for col in num_cols:
                row[col] = self._truncate4(means[col])
            summary_rows.append(row)

        summary_df = pd.DataFrame(summary_rows).sort_values(["dataset", "model", "technique"])

        reduction_rows = []
        for (ds, md), gp in summary_df.groupby(["dataset", "model"]):
            direct = gp[gp["technique"] == "Direct"]
            if not direct.empty:
                d_row = direct.iloc[0]
                for _, row in gp.iterrows():
                    res = {"dataset": ds, "model": md, "technique": row["technique"]}
                    for m in ["bandit_rate", "bandit_density", "semgrep_rate", "semgrep_density"]:
                        val = row.get(m, 0); base = d_row.get(m, 0)
                        res[m] = self._truncate4(val)
                        res[f"%▼{m}"] = self._truncate4(((base - val) / base * 100) if base else 0)
                    reduction_rows.append(res)
        reduction_df = pd.DataFrame(reduction_rows)

        loc_summary_df = pd.DataFrame(loc_rows)
        if not loc_summary_df.empty:
            loc_summary_df = loc_summary_df.groupby(["dataset", "model", "technique"])["total LOC (L)"].mean().reset_index()
            loc_summary_df["total LOC (L)"] = loc_summary_df["total LOC (L)"].apply(self._truncate4)
        
        functionality_df = self._read_functionality_data()

        with pd.ExcelWriter(self.output_file, engine="openpyxl") as writer:
            summary_df.to_excel(writer, sheet_name="summary", index=False)
            
            token_summary_df.to_excel(writer, sheet_name="Token_Usage_Sum", index=False)
            
            reduction_df.to_excel(writer, sheet_name="reduction_rate_density", index=False)
            loc_summary_df.to_excel(writer, sheet_name="Average Generated LOC", index=False)
            if not functionality_df.empty:
                functionality_df.to_excel(writer, sheet_name="functionality", index=False)

        self._format_excel()
        print(f"Success! Summary saved at: {self.output_file}")

    def _read_functionality_data(self):
        rows = []

        if not self.functionality_dir.exists():
            return pd.DataFrame()

        for root, _, files in os.walk(self.functionality_dir):
            rel_path = os.path.relpath(root, self.functionality_dir)
            parts = rel_path.split(os.sep)

            if len(parts) < 3:
                continue

            dataset = parts[0]
            model = parts[1]
            technique = self._normalize_technique(parts[2])

            for file in files:
                if not file.endswith(".txt"):
                    continue

                full_path = Path(root) / file

                try:
                    with open(full_path, "r", encoding="utf-8", errors="ignore") as f:
                        content = f.read()
                        match = self.PASS1_RE.search(content)

                        if match:
                            rows.append(
                                {
                                    "dataset": dataset,
                                    "model": model,
                                    "technique": technique,
                                    "pass@1": float(match.group(1)),
                                }
                            )
                except:
                    pass

        df = pd.DataFrame(rows)

        if df.empty:
            return df

        df = (
            df.groupby(["dataset", "model", "technique"])["pass@1"]
            .mean()
            .reset_index()
        )

        df["pass@1"] = df["pass@1"].apply(self._truncate4)

        return df


if __name__ == "__main__":
    aggregator = SecurityDataAggregator()
    aggregator.run_aggregation()